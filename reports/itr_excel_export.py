"""
ITR & CA Export Module - Excel Export Functions
openpyxl se professional Excel files generate karta hai
"""

import os
from datetime import datetime
from typing import Optional

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# ─────────────────────────────────────────────
# STYLE CONSTANTS
# ─────────────────────────────────────────────

COLOR_HEADER_BG    = "1F4E79"   # Dark blue - main header
COLOR_SUBHEAD_BG   = "2E75B6"   # Medium blue - column headers
COLOR_ALT_ROW      = "EBF3FB"   # Light blue - alternate rows
COLOR_TOTAL_BG     = "FFF2CC"   # Yellow - totals row
COLOR_WHITE        = "FFFFFF"
COLOR_FONT_LIGHT   = "FFFFFF"
COLOR_FONT_DARK    = "000000"
COLOR_POSITIVE     = "1E8449"   # Green - profit
COLOR_NEGATIVE     = "C0392B"   # Red - loss/payable

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

MEDIUM_BORDER = Border(
    left=Side(style="medium"),
    right=Side(style="medium"),
    top=Side(style="medium"),
    bottom=Side(style="medium"),
)


def _header_font(size=11, bold=True, color=COLOR_FONT_LIGHT):
    return Font(name="Calibri", size=size, bold=bold, color=color)


def _normal_font(size=10, bold=False, color=COLOR_FONT_DARK):
    return Font(name="Calibri", size=size, bold=bold, color=color)


def _fill(color):
    return PatternFill(fill_type="solid", fgColor=color)


def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


def _right():
    return Alignment(horizontal="right", vertical="center")


def _currency_fmt():
    return '₹#,##0.00'


def _write_company_header(ws, company_info: dict, title: str, fy_label: str,
                           from_date: str, to_date: str, col_span: int):
    """Sheet ke upar company header aur report title likhta hai"""
    # Row 1: Company Name
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_span)
    cell = ws.cell(row=1, column=1)
    cell.value = company_info.get("company_name", "My Company").upper()
    cell.font = Font(name="Calibri", size=16, bold=True, color=COLOR_FONT_LIGHT)
    cell.fill = _fill(COLOR_HEADER_BG)
    cell.alignment = _center()
    ws.row_dimensions[1].height = 28

    # Row 2: GST Number
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_span)
    cell = ws.cell(row=2, column=1)
    gst = company_info.get("gst_number", "")
    cell.value = f"GSTIN: {gst}" if gst else company_info.get("address", "")
    cell.font = Font(name="Calibri", size=11, bold=False, color=COLOR_FONT_LIGHT)
    cell.fill = _fill(COLOR_HEADER_BG)
    cell.alignment = _center()

    # Row 3: Report Title
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=col_span)
    cell = ws.cell(row=3, column=1)
    cell.value = title
    cell.font = Font(name="Calibri", size=13, bold=True, color=COLOR_FONT_DARK)
    cell.fill = _fill("D6E4F0")
    cell.alignment = _center()
    ws.row_dimensions[3].height = 22

    # Row 4: Financial Year & Date Range
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=col_span)
    cell = ws.cell(row=4, column=1)
    cell.value = f"Financial Year: {fy_label}   |   Period: {from_date} to {to_date}"
    cell.font = Font(name="Calibri", size=10, bold=False, color=COLOR_FONT_DARK)
    cell.fill = _fill("EBF3FB")
    cell.alignment = _center()

    # Blank spacer row
    ws.row_dimensions[5].height = 6

    return 6   # Next data starts at row 6


def _write_column_headers(ws, headers: list, row: int):
    """Column headers likhta hai with styling"""
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = _header_font(size=10)
        cell.fill = _fill(COLOR_SUBHEAD_BG)
        cell.alignment = _center()
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 20


def _auto_column_width(ws):
    """Columns ki width automatically adjust karta hai"""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val = str(cell.value) if cell.value else ""
                max_len = max(max_len, len(val))
            except Exception:
                pass
        adjusted = min(max_len + 4, 45)
        ws.column_dimensions[col_letter].width = max(adjusted, 10)


def _write_data_rows(ws, data: list, columns: list, start_row: int,
                     currency_cols: list = None, date_cols: list = None):
    """Data rows likhta hai with alternate row coloring"""
    currency_cols = currency_cols or []
    date_cols = date_cols or []

    for row_idx, row_data in enumerate(data):
        excel_row = start_row + row_idx
        bg = COLOR_ALT_ROW if row_idx % 2 == 0 else COLOR_WHITE

        for col_idx, col_key in enumerate(columns, 1):
            val = row_data.get(col_key, "")
            cell = ws.cell(row=excel_row, column=col_idx, value=val)
            cell.font = _normal_font()
            cell.fill = _fill(bg)
            cell.border = THIN_BORDER

            if col_idx in currency_cols and val != "":
                try:
                    cell.value = float(val) if val else 0.0
                    cell.number_format = _currency_fmt()
                    cell.alignment = _right()
                except Exception:
                    pass
            elif col_idx in date_cols:
                cell.alignment = _center()
            else:
                cell.alignment = _left()

        ws.row_dimensions[excel_row].height = 16

    return start_row + len(data)


def _write_totals_row(ws, data: list, columns: list, row: int,
                      label_col: int, label: str, sum_cols: list):
    """Totals row likhta hai"""
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=row, column=col_idx)
        cell.font = Font(name="Calibri", size=10, bold=True, color=COLOR_FONT_DARK)
        cell.fill = _fill(COLOR_TOTAL_BG)
        cell.border = THIN_BORDER

        if col_idx == label_col:
            cell.value = label
            cell.alignment = _left()
        elif col_idx in sum_cols:
            col_key = columns[col_idx - 1]
            total = sum(float(r.get(col_key, 0) or 0) for r in data)
            cell.value = total
            cell.number_format = _currency_fmt()
            cell.alignment = _right()
    ws.row_dimensions[row].height = 18


# ─────────────────────────────────────────────
# INDIVIDUAL SHEET WRITERS
# ─────────────────────────────────────────────

def write_sales_sheet(ws, sales_data: list, company_info: dict,
                      fy_label: str, from_date: str, to_date: str):
    headers = [
        "Sr.", "Invoice No.", "Invoice Date", "Customer Name", "Customer GST",
        "Taxable Amt (₹)", "CGST (₹)", "SGST (₹)", "IGST (₹)",
        "Total Amt (₹)", "Paid Amt (₹)", "Balance (₹)",
        "Payment Status", "Payment Method"
    ]
    col_keys = [
        None, "invoice_number", "invoice_date", "customer_name", "customer_gst",
        "taxable_amount", "cgst_amount", "sgst_amount", "igst_amount",
        "total_amount", "paid_amount", "balance_amount",
        "payment_status", "payment_method"
    ]
    currency_cols = [6, 7, 8, 9, 10, 11, 12]

    data_start = _write_company_header(ws, company_info, "SALES REGISTER",
                                        fy_label, from_date, to_date, len(headers))
    _write_column_headers(ws, headers, data_start)

    for i, row in enumerate(sales_data, 1):
        row["_sr"] = i
    col_keys[0] = "_sr"

    next_row = _write_data_rows(ws, sales_data, col_keys,
                                 data_start + 1, currency_cols, [3])

    _write_totals_row(ws, sales_data, col_keys, next_row, 4,
                      "TOTAL", currency_cols)

    ws.freeze_panes = f"A{data_start + 1}"
    _auto_column_width(ws)


def write_purchase_sheet(ws, purchase_data: list, company_info: dict,
                          fy_label: str, from_date: str, to_date: str):
    headers = [
        "Sr.", "Purchase No.", "Supplier Name", "Supplier GST",
        "Invoice No.", "Invoice Date", "Taxable Amt (₹)",
        "GST Amt (₹)", "Total Amt (₹)", "Paid Amt (₹)", "Balance (₹)"
    ]
    col_keys = [
        None, "purchase_number", "supplier_name", "supplier_gst",
        "invoice_number", "invoice_date", "taxable_amount",
        "gst_amount", "total_amount", "paid_amount", "balance_amount"
    ]
    currency_cols = [7, 8, 9, 10, 11]

    data_start = _write_company_header(ws, company_info, "PURCHASE REGISTER",
                                        fy_label, from_date, to_date, len(headers))
    _write_column_headers(ws, headers, data_start)

    for i, row in enumerate(purchase_data, 1):
        row["_sr"] = i
    col_keys[0] = "_sr"

    next_row = _write_data_rows(ws, purchase_data, col_keys,
                                 data_start + 1, currency_cols, [6])
    _write_totals_row(ws, purchase_data, col_keys, next_row, 3,
                      "TOTAL", currency_cols)

    ws.freeze_panes = f"A{data_start + 1}"
    _auto_column_width(ws)


def write_expense_sheet(ws, expense_data: list, company_info: dict,
                         fy_label: str, from_date: str, to_date: str):
    headers = [
        "Sr.", "Expense Date", "Category", "Expense Name",
        "Amount (₹)", "GST Amount (₹)", "Total Amount (₹)"
    ]
    col_keys = [
        None, "expense_date", "category", "expense_name",
        "amount", "gst_amount", "total_amount"
    ]
    currency_cols = [5, 6, 7]

    data_start = _write_company_header(ws, company_info, "EXPENSE REGISTER",
                                        fy_label, from_date, to_date, len(headers))
    _write_column_headers(ws, headers, data_start)

    for i, row in enumerate(expense_data, 1):
        row["_sr"] = i
    col_keys[0] = "_sr"

    next_row = _write_data_rows(ws, expense_data, col_keys,
                                 data_start + 1, currency_cols, [2])
    _write_totals_row(ws, expense_data, col_keys, next_row, 4,
                      "TOTAL", currency_cols)

    ws.freeze_panes = f"A{data_start + 1}"
    _auto_column_width(ws)


def write_customer_ledger_sheet(ws, customer_data: list, company_info: dict,
                                 fy_label: str, from_date: str, to_date: str):
    headers = [
        "Sr.", "Customer Name", "Customer GST", "Total Invoices",
        "Total Sales (₹)", "Total Received (₹)", "Outstanding Balance (₹)"
    ]
    col_keys = [
        None, "customer_name", "customer_gst", "total_invoices",
        "total_sales", "total_received", "outstanding_balance"
    ]
    currency_cols = [5, 6, 7]

    data_start = _write_company_header(ws, company_info, "CUSTOMER LEDGER",
                                        fy_label, from_date, to_date, len(headers))
    _write_column_headers(ws, headers, data_start)

    for i, row in enumerate(customer_data, 1):
        row["_sr"] = i
    col_keys[0] = "_sr"

    next_row = _write_data_rows(ws, customer_data, col_keys,
                                 data_start + 1, currency_cols)
    _write_totals_row(ws, customer_data, col_keys, next_row, 2,
                      "TOTAL", currency_cols)

    ws.freeze_panes = f"A{data_start + 1}"
    _auto_column_width(ws)


def write_supplier_ledger_sheet(ws, supplier_data: list, company_info: dict,
                                 fy_label: str, from_date: str, to_date: str):
    headers = [
        "Sr.", "Supplier Name", "Supplier GST", "Total Bills",
        "Total Purchase (₹)", "Total Paid (₹)", "Outstanding Balance (₹)"
    ]
    col_keys = [
        None, "supplier_name", "supplier_gst", "total_bills",
        "total_purchase", "total_paid", "outstanding_balance"
    ]
    currency_cols = [5, 6, 7]

    data_start = _write_company_header(ws, company_info, "SUPPLIER LEDGER",
                                        fy_label, from_date, to_date, len(headers))
    _write_column_headers(ws, headers, data_start)

    for i, row in enumerate(supplier_data, 1):
        row["_sr"] = i
    col_keys[0] = "_sr"

    next_row = _write_data_rows(ws, supplier_data, col_keys,
                                 data_start + 1, currency_cols)
    _write_totals_row(ws, supplier_data, col_keys, next_row, 2,
                      "TOTAL", currency_cols)

    ws.freeze_panes = f"A{data_start + 1}"
    _auto_column_width(ws)


def write_gst_summary_sheet(ws, gst_summary: dict, company_info: dict,
                              fy_label: str, from_date: str, to_date: str):
    _write_company_header(ws, company_info, "GST SUMMARY",
                           fy_label, from_date, to_date, 3)

    row = 7
    sections = [
        ("OUTPUT GST (Sales)", [
            ("Total Taxable Sales",      gst_summary.get("Total Taxable Sales",    gst_summary.get("taxable_sales", 0))),
            ("Total CGST Collected",     gst_summary.get("Total CGST",             gst_summary.get("total_cgst", 0))),
            ("Total SGST Collected",     gst_summary.get("Total SGST",             gst_summary.get("total_sgst", 0))),
            ("Total IGST Collected",     gst_summary.get("Total IGST",             gst_summary.get("total_igst", 0))),
            ("Total Output GST",         gst_summary.get("Total Output GST",       gst_summary.get("output_gst", 0))),
        ]),
        ("INPUT GST (Purchases)", [
            ("Total Taxable Purchases",  gst_summary.get("Total Taxable Purchases",gst_summary.get("taxable_purchases", 0))),
            ("Total Input GST",          gst_summary.get("Total Input GST",        gst_summary.get("input_gst", 0))),
        ]),
        ("NET GST PAYABLE", [
            ("Net GST Payable (Output - Input)", gst_summary.get("GST Payable (Net)", gst_summary.get("net_gst_payable", 0))),
        ]),
    ]

    for section_title, items in sections:
        # Section title
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        cell = ws.cell(row=row, column=1, value=section_title)
        cell.font = _header_font(size=11)
        cell.fill = _fill(COLOR_SUBHEAD_BG)
        cell.alignment = _center()
        cell.border = THIN_BORDER
        ws.row_dimensions[row].height = 20
        row += 1

        for label, value in items:
            ws.cell(row=row, column=1, value="").fill = _fill(COLOR_WHITE)
            lbl_cell = ws.cell(row=row, column=2, value=label)
            lbl_cell.font = _normal_font(bold=True)
            lbl_cell.alignment = _left()
            lbl_cell.border = THIN_BORDER

            val_cell = ws.cell(row=row, column=3, value=float(value or 0))
            val_cell.number_format = _currency_fmt()
            val_cell.alignment = _right()
            val_cell.border = THIN_BORDER

            if section_title == "NET GST PAYABLE":
                color = COLOR_NEGATIVE if value > 0 else COLOR_POSITIVE
                val_cell.font = Font(name="Calibri", size=11, bold=True, color=color)
                lbl_cell.fill = _fill(COLOR_TOTAL_BG)
                val_cell.fill = _fill(COLOR_TOTAL_BG)
            else:
                lbl_cell.fill = _fill(COLOR_ALT_ROW if row % 2 == 0 else COLOR_WHITE)
                val_cell.fill = _fill(COLOR_ALT_ROW if row % 2 == 0 else COLOR_WHITE)

            ws.row_dimensions[row].height = 18
            row += 1
        row += 1  # gap between sections

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 22


def write_pl_summary_sheet(ws, pl_summary: dict, company_info: dict,
                            fy_label: str, from_date: str, to_date: str):
    _write_company_header(ws, company_info, "PROFIT & LOSS SUMMARY",
                           fy_label, from_date, to_date, 3)

    row = 7
    pl_items = [
        ("INCOME", [
            ("Total Sales (Incl. GST)",   pl_summary.get("Total Sales",     pl_summary.get("total_sales", 0)),     False),
            ("Taxable Sales",              pl_summary.get("Taxable Sales",   pl_summary.get("taxable_sales", 0)),   False),
        ]),
        ("EXPENSES", [
            ("Total Purchases",            pl_summary.get("Total Purchases", pl_summary.get("total_purchases", 0)), False),
            ("Total Expenses",             pl_summary.get("Total Expenses",  pl_summary.get("total_expenses", 0)),  False),
        ]),
        ("PROFIT / LOSS", [
            ("Gross Profit (Sales - Purchases)", pl_summary.get("Gross Profit", pl_summary.get("gross_profit", 0)), True),
            ("Net Profit (Gross - Expenses)",    pl_summary.get("Net Profit",   pl_summary.get("net_profit", 0)),   True),
        ]),
    ]

    for section_title, items in pl_items:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        cell = ws.cell(row=row, column=1, value=section_title)
        cell.font = _header_font(size=11)
        cell.fill = _fill(COLOR_SUBHEAD_BG)
        cell.alignment = _center()
        cell.border = THIN_BORDER
        ws.row_dimensions[row].height = 20
        row += 1

        for label, value, is_result in items:
            lbl_cell = ws.cell(row=row, column=2, value=label)
            val_cell = ws.cell(row=row, column=3, value=float(value or 0))
            val_cell.number_format = _currency_fmt()
            val_cell.alignment = _right()
            lbl_cell.border = THIN_BORDER
            val_cell.border = THIN_BORDER
            ws.cell(row=row, column=1).fill = _fill(COLOR_WHITE)

            if is_result:
                color = COLOR_POSITIVE if value >= 0 else COLOR_NEGATIVE
                lbl_cell.font = Font(name="Calibri", size=11, bold=True, color=COLOR_FONT_DARK)
                val_cell.font = Font(name="Calibri", size=11, bold=True, color=color)
                lbl_cell.fill = _fill(COLOR_TOTAL_BG)
                val_cell.fill = _fill(COLOR_TOTAL_BG)
                ws.row_dimensions[row].height = 20
            else:
                lbl_cell.font = _normal_font(bold=False)
                val_cell.font = _normal_font()
                bg = COLOR_ALT_ROW if row % 2 == 0 else COLOR_WHITE
                lbl_cell.fill = _fill(bg)
                val_cell.fill = _fill(bg)
                ws.row_dimensions[row].height = 18
            row += 1
        row += 1

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 22


# ─────────────────────────────────────────────
# MAIN EXPORT FUNCTIONS
# ─────────────────────────────────────────────

def export_single_excel(data: list, sheet_writer_func, company_info: dict,
                         fy_label: str, from_date: str, to_date: str,
                         sheet_name: str, save_path: str) -> bool:
    """Single sheet Excel export"""
    if not OPENPYXL_AVAILABLE:
        raise ImportError("openpyxl install nahi hai. 'pip install openpyxl' chalayein.")

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.sheet_view.showGridLines = False

    sheet_writer_func(ws, data, company_info, fy_label, from_date, to_date)

    wb.save(save_path)
    return True


def export_complete_itr(
    sales_data, purchase_data, expense_data,
    customer_data, supplier_data,
    gst_summary, pl_summary,
    company_info,           # dict ya string dono accept karta hai
    from_date: str, to_date: str,
    save_path: str,
    fy_label: str = None    # optional — auto calculate ho jaata hai
) -> bool:
    """Complete ITR Excel file - multiple sheets ke saath"""
    if not OPENPYXL_AVAILABLE:
        raise ImportError("openpyxl install nahi hai. 'pip install openpyxl' chalayein.")

    # company_info string bhi ho sakta hai (itr_export_ui.py se)
    if isinstance(company_info, str):
        company_info = {"company_name": company_info}

    # fy_label auto calculate
    if not fy_label:
        fy_start = int(from_date[:4])
        fy_label = f"{fy_start}-{str(fy_start + 1)[-2:]}"

    wb = Workbook()

    sheet_configs = [
        ("Sales Register",     write_sales_sheet,           sales_data),
        ("Purchase Register",  write_purchase_sheet,        purchase_data),
        ("Expense Register",   write_expense_sheet,         expense_data),
        ("Customer Ledger",    write_customer_ledger_sheet, customer_data),
        ("Supplier Ledger",    write_supplier_ledger_sheet, supplier_data),
    ]

    first = True
    for sheet_name, writer_func, data in sheet_configs:
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(title=sheet_name)

        ws.sheet_view.showGridLines = False
        writer_func(ws, data, company_info, fy_label, from_date, to_date)

    # GST Summary sheet
    ws_gst = wb.create_sheet(title="GST Summary")
    ws_gst.sheet_view.showGridLines = False
    write_gst_summary_sheet(ws_gst, gst_summary, company_info, fy_label, from_date, to_date)

    # P&L Summary sheet
    ws_pl = wb.create_sheet(title="Profit & Loss")
    ws_pl.sheet_view.showGridLines = False
    write_pl_summary_sheet(ws_pl, pl_summary, company_info, fy_label, from_date, to_date)

    wb.save(save_path)
    return True


def generate_filename(from_date: str, to_date: str, prefix: str = "ITR_Report") -> str:
    """Export file ka naam generate karta hai"""
    from_clean = from_date.replace("-", "_")
    to_clean = to_date.replace("-", "_")
    return f"{prefix}_{from_clean}_To_{to_clean}.xlsx"


# ─────────────────────────────────────────────
# ALIAS FUNCTIONS (itr_export_ui.py compatibility)
# ─────────────────────────────────────────────

# itr_export_ui.py mein build_file_name naam se call hota hai
def build_file_name(from_date: str, to_date: str, prefix: str = "ITR_Report") -> str:
    return generate_filename(from_date, to_date, prefix)


# itr_export_ui.py mein get_default_save_dir naam se call hota hai
def get_default_save_dir(location_key: str = "desktop") -> str:
    home = os.path.expanduser("~")
    mapping = {
        "desktop":   os.path.join(home, "Desktop"),
        "documents": os.path.join(home, "Documents"),
        "downloads": os.path.join(home, "Downloads"),
    }
    path = mapping.get(location_key.lower(), os.path.join(home, "Desktop"))
    os.makedirs(path, exist_ok=True)
    return path


# itr_export_ui.py mein export_single_sheet naam se call hota hai
_SHEET_WRITER_MAP = {
    "sales":           (write_sales_sheet,           "Sales Register"),
    "purchase":        (write_purchase_sheet,        "Purchase Register"),
    "expense":         (write_expense_sheet,         "Expense Register"),
    "customer_ledger": (write_customer_ledger_sheet, "Customer Ledger"),
    "supplier_ledger": (write_supplier_ledger_sheet, "Supplier Ledger"),
}

def export_single_sheet(data: list, sheet_type: str, company_name: str,
                         from_date: str, to_date: str, save_path: str) -> bool:
    """Single sheet export — itr_export_ui.py ke saath compatible"""
    writer_func, sheet_name = _SHEET_WRITER_MAP[sheet_type]
    company_info = {"company_name": company_name}
    fy_start = int(from_date[:4])
    fy_label = f"{fy_start}-{str(fy_start + 1)[-2:]}"
    return export_single_excel(
        data, writer_func, company_info,
        fy_label, from_date, to_date,
        sheet_name, save_path
    )
