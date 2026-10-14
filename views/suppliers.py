from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel,
    QPushButton, QLineEdit, QTableWidget, QTableWidgetItem, QDialog,
    QFormLayout, QMessageBox, QHeaderView, QAbstractItemView, QComboBox,
    QDateEdit, QFileDialog, QFrame, QTabWidget)
from PyQt5.QtCore import Qt, QDate
from PyQt5.QtGui import QFont, QColor
import os, sqlite3


# ── Supplier Add/Edit Dialog ──────────────────────────────────────────────────
class SupplierDialog(QDialog):
    def __init__(self, db, supplier=None, parent=None):
        super().__init__(parent)
        self.db = db
        self.setWindowTitle("Edit Supplier" if supplier else "New Supplier")
        self.setMinimumWidth(480)
        layout = QVBoxLayout(self)
        form = QFormLayout()
        self.fields = {}
        for key, label in [
            ("supplier_name","Name *"), ("company_name","Company"),
            ("mobile","Mobile"),       ("email","Email"),
            ("gst_number","GST Number"),("pan_number","PAN"),
            ("city","City"),           ("state","State"),
        ]:
            w = QLineEdit(); w.setFixedHeight(36)
            form.addRow(label+":", w); self.fields[key] = w
        self.status = QComboBox(); self.status.setFixedHeight(36)
        self.status.addItems(["active","inactive","blocked","preferred"])
        form.addRow("Status:", self.status)
        layout.addLayout(form)
        row = QHBoxLayout()
        s = QPushButton("💾 Save"); s.setFixedHeight(40)
        s.setStyleSheet("background:#27AE60;color:white;border-radius:6px;font-weight:700;")
        s.clicked.connect(self.accept)
        c = QPushButton("Cancel"); c.setFixedHeight(40)
        c.setStyleSheet("background:#E74C3C;color:white;border-radius:6px;font-weight:700;")
        c.clicked.connect(self.reject)
        row.addWidget(s); row.addWidget(c); layout.addLayout(row)
        if supplier:
            for k, w in self.fields.items(): w.setText(str(supplier.get(k,"") or ""))
            self.status.setCurrentText(supplier.get("status","active"))

    def get_data(self):
        d = {k: w.text().strip() for k, w in self.fields.items()}
        d["status"] = self.status.currentText()
        return d


# ── Supplier Ledger Dialog ────────────────────────────────────────────────────
class SupplierLedgerDialog(QDialog):
    def __init__(self, db, supplier_id, supplier_name, parent=None):
        super().__init__(parent)
        self.db          = db
        self.supplier_id = supplier_id
        self.sup_name    = supplier_name
        self.setWindowTitle(f"Supplier Ledger — {supplier_name}")
        self.setMinimumSize(820, 560)

        layout = QVBoxLayout(self)
        layout.setSpacing(10)
        layout.setContentsMargins(14, 14, 14, 14)

        # ── Title ─────────────────────────────────────────────────────────
        title = QLabel(f"📋  {supplier_name} — Ledger")
        title.setFont(QFont("Segoe UI", 14, 700))
        title.setStyleSheet("color:#1A2A4A;")
        layout.addWidget(title)

        # ── Date range row ────────────────────────────────────────────────
        dr = QHBoxLayout()
        dr.addWidget(QLabel("From:"))
        self.from_date = QDateEdit()
        self.from_date.setCalendarPopup(True)
        self.from_date.setDate(QDate.currentDate().addMonths(-1))
        self.from_date.setDisplayFormat("dd-MM-yyyy")
        self.from_date.setFixedHeight(34)

        dr.addWidget(self.from_date)
        dr.addSpacing(10)
        dr.addWidget(QLabel("To:"))
        self.to_date = QDateEdit()
        self.to_date.setCalendarPopup(True)
        self.to_date.setDate(QDate.currentDate())
        self.to_date.setDisplayFormat("dd-MM-yyyy")
        self.to_date.setFixedHeight(34)
        dr.addWidget(self.to_date)
        dr.addSpacing(14)

        load_btn = QPushButton("🔍 Load")
        load_btn.setFixedHeight(34)
        load_btn.setStyleSheet("background:#2980B9;color:white;border-radius:6px;font-weight:600;padding:0 14px;")
        load_btn.clicked.connect(self._load)
        dr.addWidget(load_btn)

        dl_btn = QPushButton("⬇ Download Excel")
        dl_btn.setFixedHeight(34)
        dl_btn.setStyleSheet("background:#27AE60;color:white;border-radius:6px;font-weight:700;padding:0 14px;")
        dl_btn.clicked.connect(self._download)
        dr.addWidget(dl_btn)
        dr.addStretch()
        layout.addLayout(dr)

        # ── Tabs: Ledger | Purchases ──────────────────────────────────────
        self.tabs = QTabWidget()
        layout.addWidget(self.tabs)

        # --- Ledger tab ---
        ledger_tab = QWidget()
        lt = QVBoxLayout(ledger_tab)
        lt.setContentsMargins(0, 8, 0, 0)
        self.ledger_table = QTableWidget()
        self.ledger_table.setColumnCount(6)
        self.ledger_table.setHorizontalHeaderLabels(
            ["Date","Description","Debit (₹)","Credit (₹)","Balance (₹)","Purchase #"])
        self.ledger_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        self.ledger_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.ledger_table.setAlternatingRowColors(True)
        self.ledger_table.verticalHeader().setVisible(False)
        lt.addWidget(self.ledger_table)
        self.tabs.addTab(ledger_tab, "📒 Ledger")

        # --- Purchases tab ---
        pur_tab = QWidget()
        pt = QVBoxLayout(pur_tab)
        pt.setContentsMargins(0, 8, 0, 0)
        self.pur_table = QTableWidget()
        self.pur_table.setColumnCount(6)
        self.pur_table.setHorizontalHeaderLabels(
            ["Purchase #","Date","Total (₹)","Paid (₹)","Balance (₹)","Status"])
        self.pur_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.pur_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.pur_table.setAlternatingRowColors(True)
        self.pur_table.verticalHeader().setVisible(False)
        pt.addWidget(self.pur_table)
        self.tabs.addTab(pur_tab, "🛒 Purchases")

        # ── Summary bar ───────────────────────────────────────────────────
        self.summary_lbl = QLabel("")
        self.summary_lbl.setFont(QFont("Segoe UI", 11, 700))
        self.summary_lbl.setStyleSheet(
            "color:#1A2A4A; background:#EEF2FF; border-radius:6px; padding:6px 12px;")
        layout.addWidget(self.summary_lbl)

        # ── Close ─────────────────────────────────────────────────────────
        close_btn = QPushButton("✖ Close")
        close_btn.setFixedHeight(36)
        close_btn.setStyleSheet("background:#1A2A4A;color:white;border-radius:6px;font-weight:600;")
        close_btn.clicked.connect(self.accept)
        layout.addWidget(close_btn)

        self._load()

    # ── Fetch helpers ─────────────────────────────────────────────────────
    def _dates(self):
        return (self.from_date.date().toString("yyyy-MM-dd"),
                self.to_date.date().toString("yyyy-MM-dd"))

    def _fetch_ledger(self):
        fd, td = self._dates()
        try:
            conn = sqlite3.connect(self.db.db_path)
            conn.row_factory = sqlite3.Row
            rows = conn.execute("""
                SELECT sl.entry_date, sl.description,
                       sl.debit, sl.credit, sl.balance,
                       p.purchase_number
                FROM supplier_ledger sl
                LEFT JOIN purchases p ON sl.purchase_id = p.id
                WHERE sl.supplier_id = ?
                  AND sl.entry_date BETWEEN ? AND ?
                ORDER BY sl.entry_date ASC, sl.id ASC
            """, (self.supplier_id, fd, td)).fetchall()
            conn.close()
            return [dict(r) for r in rows]
        except Exception as e:
            return []

    def _fetch_purchases(self):
        fd, td = self._dates()
        try:
            conn = sqlite3.connect(self.db.db_path)
            conn.row_factory = sqlite3.Row
            rows = conn.execute("""
                SELECT purchase_number, purchase_date,
                       grand_total, paid_amount, balance_amount, payment_status
                FROM purchases
                WHERE supplier_id = ?
                  AND purchase_date BETWEEN ? AND ?
                ORDER BY purchase_date ASC
            """, (self.supplier_id, fd, td)).fetchall()
            conn.close()
            return [dict(r) for r in rows]
        except Exception as e:
            return []

    def _load(self):
        # Ledger tab
        entries = self._fetch_ledger()
        self.ledger_table.setRowCount(len(entries))
        for i, e in enumerate(entries):
            self.ledger_table.setItem(i,0,QTableWidgetItem(str(e.get("entry_date","") or "")))
            self.ledger_table.setItem(i,1,QTableWidgetItem(str(e.get("description","") or "")))
            dr = QTableWidgetItem(f"₹{e.get('debit',0):,.2f}")
            dr.setTextAlignment(Qt.AlignRight|Qt.AlignVCenter)
            if (e.get('debit') or 0) > 0: dr.setForeground(QColor("#E74C3C"))
            cr = QTableWidgetItem(f"₹{e.get('credit',0):,.2f}")
            cr.setTextAlignment(Qt.AlignRight|Qt.AlignVCenter)
            if (e.get('credit') or 0) > 0: cr.setForeground(QColor("#27AE60"))
            bal = QTableWidgetItem(f"₹{e.get('balance',0):,.2f}")
            bal.setTextAlignment(Qt.AlignRight|Qt.AlignVCenter)
            self.ledger_table.setItem(i,2,dr)
            self.ledger_table.setItem(i,3,cr)
            self.ledger_table.setItem(i,4,bal)
            self.ledger_table.setItem(i,5,QTableWidgetItem(str(e.get("purchase_number","") or "")))

        # Purchases tab
        purchases = self._fetch_purchases()
        self.pur_table.setRowCount(len(purchases))
        total_pur = 0; total_paid = 0; total_bal = 0
        for i, p in enumerate(purchases):
            gt  = p.get("grand_total",0) or 0
            pd  = p.get("paid_amount",0) or 0
            bl  = p.get("balance_amount",0) or 0
            total_pur += gt; total_paid += pd; total_bal += bl
            self.pur_table.setItem(i,0,QTableWidgetItem(str(p.get("purchase_number","") or "")))
            self.pur_table.setItem(i,1,QTableWidgetItem(str(p.get("purchase_date","") or "")))
            for col, val in [(2,gt),(3,pd),(4,bl)]:
                it = QTableWidgetItem(f"₹{val:,.2f}")
                it.setTextAlignment(Qt.AlignRight|Qt.AlignVCenter)
                self.pur_table.setItem(i,col,it)
            self.pur_table.setItem(i,5,QTableWidgetItem(
                (p.get("payment_status","") or "").upper()))

        bal_final = entries[-1]["balance"] if entries else total_bal
        self.summary_lbl.setText(
            f"Total Purchase: ₹{total_pur:,.2f}   |   "
            f"Total Paid: ₹{total_paid:,.2f}   |   "
            f"Outstanding: ₹{bal_final:,.2f}"
        )

    # ── Excel Download ────────────────────────────────────────────────────
    def _download(self):
        entries   = self._fetch_ledger()
        purchases = self._fetch_purchases()
        if not entries and not purchases:
            QMessageBox.information(self,"No Data","Is date range mein koi data nahi hai!")
            return
        fd = self.from_date.date().toString("dd-MM-yyyy")
        td = self.to_date.date().toString("dd-MM-yyyy")
        safe = self.sup_name.replace(" ","_").replace("/","-")
        default = f"Supplier_Ledger_{safe}_{fd}_to_{td}.xlsx"
        path, _ = QFileDialog.getSaveFileName(
            self, "Save Excel",
            os.path.join(os.path.expanduser("~"), "Desktop", default),
            "Excel Files (*.xlsx)"
        )
        if not path: return
        try:
            self._write_excel(entries, purchases, fd, td, path)
            QMessageBox.information(self,"✅ Downloaded", f"Saved:\n{path}")
        except Exception as e:
            QMessageBox.critical(self,"Error", str(e))

    def _write_excel(self, entries, purchases, fd, td, path):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ImportError("pip install openpyxl")

        wb  = Workbook()
        thin   = Side(style="thin",   color="000000")
        medium = Side(style="medium", color="000000")
        def bdr(thick=False):
            s = medium if thick else thin
            return Border(left=s, right=s, top=s, bottom=s)
        def hdr_font(): return Font(name="Calibri", size=11, bold=True)
        def data_font(bold=False): return Font(name="Calibri", size=10, bold=bold)

        # ── Sheet 1: Ledger ───────────────────────────────────────────────
        ws = wb.active; ws.title = "Supplier Ledger"
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:F1")
        ws["A1"] = f"Hari Om Electric — Supplier Ledger"
        ws["A1"].font      = Font(name="Calibri", size=13, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws["A1"].border    = bdr(True)
        ws.row_dimensions[1].height = 26

        ws.merge_cells("A2:F2")
        ws["A2"] = f"Supplier: {self.sup_name}   |   Period: {fd} to {td}"
        ws["A2"].font      = Font(name="Calibri", size=11, bold=True)
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws["A2"].border    = bdr()
        ws.row_dimensions[2].height = 20
        ws.append([])

        hdrs = ["Date","Description","Debit (Rs.)","Credit (Rs.)","Balance (Rs.)","Purchase #"]
        widths = [12, 35, 13, 13, 14, 13]
        for c,(h,w) in enumerate(zip(hdrs,widths),1):
            cell = ws.cell(row=4, column=c, value=h)
            cell.font = hdr_font(); cell.border = bdr(True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.row_dimensions[4].height = 18

        for idx, e in enumerate(entries):
            r = 5+idx
            row_data = [
                str(e.get("entry_date","") or ""),
                str(e.get("description","") or ""),
                e.get("debit",0) or 0,
                e.get("credit",0) or 0,
                e.get("balance",0) or 0,
                str(e.get("purchase_number","") or ""),
            ]
            for c, val in enumerate(row_data, 1):
                cell = ws.cell(row=r, column=c, value=val)
                cell.border = bdr()
                cell.font   = data_font(c in (3,4,5))
                if c in (3,4,5):
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[r].height = 16

        # Summary
        sr = 5+len(entries)
        ws.merge_cells(f"A{sr}:D{sr}")
        ws[f"A{sr}"] = "Outstanding Balance"
        ws[f"A{sr}"].font = Font(name="Calibri",size=11,bold=True)
        ws[f"A{sr}"].alignment = Alignment(horizontal="right")
        ws[f"A{sr}"].border = bdr(True)
        bal = entries[-1]["balance"] if entries else 0
        ws[f"E{sr}"] = bal
        ws[f"E{sr}"].font = Font(name="Calibri",size=12,bold=True)
        ws[f"E{sr}"].number_format = '#,##0.00'
        ws[f"E{sr}"].alignment = Alignment(horizontal="right")
        ws[f"E{sr}"].border = bdr(True)
        ws[f"F{sr}"].border = bdr(True)
        ws.row_dimensions[sr].height = 20

        # Print settings
        from openpyxl.worksheet.page import PageMargins
        ws.page_setup.paperSize   = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToPage   = True
        ws.page_setup.fitToWidth  = 1
        ws.print_title_rows = "1:4"
        ws.page_margins = PageMargins(left=0.5,right=0.5,top=0.6,bottom=0.6)
        ws.freeze_panes = "A5"

        # ── Sheet 2: Purchases ────────────────────────────────────────────
        ws2 = wb.create_sheet(title="Purchases")
        ws2.sheet_view.showGridLines = False

        ws2.merge_cells("A1:F1")
        ws2["A1"] = f"Purchases — {self.sup_name}   |   {fd} to {td}"
        ws2["A1"].font = Font(name="Calibri",size=13,bold=True)
        ws2["A1"].alignment = Alignment(horizontal="center",vertical="center")
        ws2["A1"].border = bdr(True)
        ws2.row_dimensions[1].height = 26

        p_hdrs = ["Purchase #","Date","Total (Rs.)","Paid (Rs.)","Balance (Rs.)","Status"]
        p_wids = [14, 12, 14, 14, 14, 12]
        for c,(h,w) in enumerate(zip(p_hdrs,p_wids),1):
            cell = ws2.cell(row=2, column=c, value=h)
            cell.font = hdr_font(); cell.border = bdr(True)
            cell.alignment = Alignment(horizontal="center",vertical="center")
            ws2.column_dimensions[get_column_letter(c)].width = w
        ws2.row_dimensions[2].height = 18

        tot_pur=0; tot_paid=0; tot_bal=0
        for idx, p in enumerate(purchases):
            r = 3+idx
            gt = p.get("grand_total",0) or 0
            pd = p.get("paid_amount",0) or 0
            bl = p.get("balance_amount",0) or 0
            tot_pur+=gt; tot_paid+=pd; tot_bal+=bl
            row_data = [
                str(p.get("purchase_number","") or ""),
                str(p.get("purchase_date","") or ""),
                gt, pd, bl,
                (p.get("payment_status","") or "").upper()
            ]
            for c, val in enumerate(row_data,1):
                cell = ws2.cell(row=r, column=c, value=val)
                cell.border = bdr()
                cell.font = data_font(c in (3,4,5))
                if c in (3,4,5):
                    cell.number_format='#,##0.00'
                    cell.alignment=Alignment(horizontal="right",vertical="center")
                else:
                    cell.alignment=Alignment(horizontal="left",vertical="center")
            ws2.row_dimensions[r].height = 16

        # Total row
        tr = 3+len(purchases)
        ws2.cell(row=tr,column=1,value="TOTAL").font = Font(name="Calibri",size=11,bold=True)
        ws2.cell(row=tr,column=1).border=bdr(True)
        ws2.cell(row=tr,column=2).border=bdr(True)
        for c, val in [(3,tot_pur),(4,tot_paid),(5,tot_bal)]:
            cell = ws2.cell(row=tr, column=c, value=val)
            cell.font = Font(name="Calibri",size=11,bold=True)
            cell.number_format='#,##0.00'
            cell.alignment=Alignment(horizontal="right")
            cell.border=bdr(True)
        ws2.cell(row=tr,column=6).border=bdr(True)
        ws2.row_dimensions[tr].height=20

        ws2.page_setup.paperSize  = ws2.PAPERSIZE_A4
        ws2.page_setup.fitToPage  = True
        ws2.page_setup.fitToWidth = 1
        ws2.freeze_panes = "A3"

        wb.save(path)


# ── Main Suppliers Widget ─────────────────────────────────────────────────────
class SuppliersWidget(QWidget):
    def __init__(self, db, user):
        super().__init__()
        self.db = db; self.user = user
        layout = QVBoxLayout(self)
        layout.setContentsMargins(16,16,16,16); layout.setSpacing(12)

        hr = QHBoxLayout()
        t = QLabel("Supplier Management")
        t.setFont(QFont("Segoe UI",16,700)); t.setStyleSheet("color:#1A2A4A;")
        b = QPushButton("+ New Supplier"); b.setFixedHeight(38)
        b.setStyleSheet("background:#1A2A4A;color:white;border-radius:6px;font-weight:600;padding:0 14px;")
        b.clicked.connect(self._add)
        hr.addWidget(t); hr.addStretch(); hr.addWidget(b)
        layout.addLayout(hr)

        self.search = QLineEdit()
        self.search.setPlaceholderText("🔍 Search suppliers...")
        self.search.setFixedHeight(38)
        self.search.textChanged.connect(self._load)
        layout.addWidget(self.search)

        self.table = QTableWidget(); self.table.setColumnCount(8)
        self.table.setHorizontalHeaderLabels(
            ["Name","Company","Mobile","GST","City","State","Status","Action"])
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setVisible(False)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setColumnWidth(7, 130)
        self.table.doubleClicked.connect(self._edit)
        layout.addWidget(self.table)

        ar = QHBoxLayout()
        for lbl, fn, style in [
            ("✏️ Edit",   self._edit,   "background:#F39C12;color:white;border-radius:6px;font-weight:600;padding:0 14px;"),
            ("🗑 Delete", self._delete, "background:#E74C3C;color:white;border-radius:6px;font-weight:600;padding:0 14px;"),
        ]:
            btn = QPushButton(lbl); btn.setFixedHeight(36)
            btn.setStyleSheet(style); btn.clicked.connect(fn)
            ar.addWidget(btn)
        ar.addStretch(); layout.addLayout(ar)
        self._load()

    def refresh(self): self._load()

    def _load(self):
        rows = self.db.get_suppliers(search=self.search.text())
        self.table.setRowCount(len(rows))
        for i, s in enumerate(rows):
            for j, key in enumerate(["supplier_name","company_name","mobile",
                                      "gst_number","city","state","status"]):
                item = QTableWidgetItem(str(s.get(key,"") or ""))
                item.setData(Qt.UserRole, s["id"])
                self.table.setItem(i, j, item)

            # Ledger/View button
            view_btn = QPushButton("📒 Ledger")
            view_btn.setFixedHeight(28)
            view_btn.setStyleSheet(
                "background:#2980B9;color:white;border-radius:5px;font-weight:600;font-size:11px;")
            view_btn.clicked.connect(
                lambda _, sid=s["id"], sname=s.get("supplier_name",""):
                    self._open_ledger(sid, sname))
            self.table.setCellWidget(i, 7, view_btn)

    def _sel_id(self):
        r = self.table.currentRow()
        return self.table.item(r,0).data(Qt.UserRole) \
               if r>=0 and self.table.item(r,0) else None

    def _open_ledger(self, supplier_id, supplier_name):
        dlg = SupplierLedgerDialog(self.db, supplier_id, supplier_name, self)
        dlg.exec_()

    def _add(self):
        d = SupplierDialog(self.db, parent=self)
        if d.exec_():
            self.db.save_supplier(d.get_data()); self._load()

    def _edit(self):
        sid = self._sel_id()
        if not sid: return
        s = self.db.get_supplier(sid)
        d = SupplierDialog(self.db, s, self)
        if d.exec_():
            self.db.save_supplier(d.get_data(), sid); self._load()

    def _delete(self):
        if self.user["role"] != "admin":
            QMessageBox.warning(self,"Permission","Admin only."); return
        sid = self._sel_id()
        if sid and QMessageBox.question(
            self,"Delete","Delete supplier?",
            QMessageBox.Yes|QMessageBox.No) == QMessageBox.Yes:
            c = self.db.get_connection()
            c.execute("UPDATE suppliers SET status='inactive' WHERE id=?", (sid,))
            c.commit(); c.close(); self._load()